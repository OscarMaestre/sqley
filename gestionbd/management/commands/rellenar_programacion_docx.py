from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from gestionbd.models import Modulo,ResultadoDeAprendizaje,Contenido, ResultadoContenido

from copy import deepcopy

from docx import Document
from docx.shared import RGBColor

from math import ceil


def get_run(documento, marcador):
    # Buscar en párrafos normales
    for parrafo in documento.paragraphs:
        if marcador in parrafo.text:
            # Borrar texto del párrafo
            for run in parrafo.runs:
                run.text = ""
            # Devolver primer run vacío para edición posterior
            return parrafo.runs[0]
    
    # Buscar en tablas
    for tabla in documento.tables:
        for fila in tabla.rows:
            for celda in fila.cells:
                for parrafo in celda.paragraphs:
                    if marcador in parrafo.text:
                        for run in parrafo.runs:
                            run.text = ""
                        return parrafo.runs[0]
    
    return None  # Si no se encuentra marcador

class Command(BaseCommand):
    help = "Rellena los ficheros y presentaciones de Word."
    
    def add_arguments(self, parser):
        parser.add_argument("codigo_jccm", nargs="+", type=str, help="Codigo tal y como aparece en los DOCM de la Junta, p.ej 0370 para 'Redes' 0378 para 'Seguridad y alta disponibilidad' o 0373 para 'Lenguajes de marcas'")
        parser.add_argument("archivo_programacion_docx", nargs="+", type=str, help="Ruta al archivo DOCX con la PROGRAMACION")
       
        
        parser.add_argument("archivo_resultado_programacion_docx", nargs="+", type=str, help="Ruta al archivo DOCX donde se guardará la PROGRAMACION RESULTADO")
     

    
    def _anadir_modulo_a_portada(self, doc, modulo):
        run=get_run(doc, "{{NOMBREMODULO}}")
        run.text=modulo.nombre
    def _anadir_nombre_profesor(self, doc,modulo):
        run=get_run(doc, "{{NOMBREPROFESOR}}")
        run.text="Óscar Gómez"
    def __anadir_horas_semanales(self, doc, modulo):
        run=get_run(doc, "{{NUMHORASSEMANALES}}")
        run.text=str(modulo.horas_semanales)
    def __anadir_horas_anuales(self, doc, modulo):
        run=get_run(doc, "{{NUMHORASANUALES}}")
        run.text=str(modulo.horas_anuales)
    def __anadir_horas_perdida_eval_continua(self, doc, modulo):
        PORCENTAJE_PERDIDA_EVAL_CONTINUA=0.25
        run=get_run(doc, "{{NUMHORASPERDIDAEVALUACIONCONTINUA}}")
        horas_para_perder_eval_continua=float(modulo.horas_anuales)*PORCENTAJE_PERDIDA_EVAL_CONTINUA
        horas_redondeadas=ceil(horas_para_perder_eval_continua)
        run.text=str(horas_redondeadas)
        
    def hacer_todos_reemplazos(self, doc, modulo):
        reemplazos=[
            self._anadir_modulo_a_portada,
            self._anadir_nombre_profesor,
            self.__anadir_horas_anuales,
            self.__anadir_horas_semanales,
            self.__anadir_horas_perdida_eval_continua
        ]
        for reemplazo in reemplazos:
            reemplazo(doc, modulo)
        
    def handle(self, *args, **options):
        
        codigo_junta=options["codigo_jccm"][0]
        print("Código leído: "+str(codigo_junta))
        archivo_programacion_docx=options["archivo_programacion_docx"][0]
        archivo_resultado_programacion_docx=options["archivo_resultado_programacion_docx"][0]
        modulo=Modulo.objects.filter(codigo_junta=codigo_junta)[0]
        
        print("Rellenando la programación para el módulo "+str(modulo))
        print("Abriendo el fichero:"+archivo_programacion_docx)
        doc = Document(archivo_programacion_docx)
        #reemplazamos
        self.hacer_todos_reemplazos(doc, modulo)
        #Fin del reemplazado
        print("Guardando en "+archivo_resultado_programacion_docx+"...")
        doc.save(archivo_resultado_programacion_docx)
        print("Guardado en  "+archivo_resultado_programacion_docx)
